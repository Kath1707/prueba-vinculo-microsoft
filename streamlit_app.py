"""
App de PRUEBA de vinculación con OneDrive (Microsoft Graph)
============================================================
Sirve para comprobar, sin tocar tus apps reales, que:
  1) los Secrets están bien puestos,
  2) Azure entrega un token y tiene los permisos aprobados,
  3) la app ve tu OneDrive y la carpeta de trabajo,
  4) puede crear, leer y borrar archivos,
  5) puede editar un Excel (copiar hoja plantilla, insertar filas antes de la firma) y subirlo
     de vuelta sin pisar cambios de otra persona.

Solo crea archivos que empiezan con "_PRUEBA_VINCULACION" y los borra al terminar.
NUNCA muestra ni guarda secretos.

Secrets (Streamlit Cloud -> Manage app -> Settings -> Secrets), NO en GitHub:

    AZURE_CLIENT_ID     = "..."
    AZURE_TENANT_ID     = "..."
    AZURE_CLIENT_SECRET = "..."   # el VALOR del secreto (no el "Id del secreto")
    ONEDRIVE_USER       = "correo@tudominio.pe"   # dueño del OneDrive
    ONEDRIVE_FOLDER     = "Calidad/Pruebas"       # carpeta dentro de ese OneDrive
    # Opcional: PIN para abrir esta app de prueba
    APP_PIN             = "1234"
"""

import base64
import io
import json
import re
import time
from datetime import datetime

import requests
import streamlit as st

try:
    import msal
    MSAL_OK = True
except ImportError:  # pragma: no cover
    MSAL_OK = False

st.set_page_config(page_title="Prueba de vinculación OneDrive", page_icon="🔗", layout="centered")

GRAPH = "https://graph.microsoft.com/v1.0"
PREFIJO = "_PRUEBA_VINCULACION"
CLAVES = ["AZURE_CLIENT_ID", "AZURE_TENANT_ID", "AZURE_CLIENT_SECRET", "ONEDRIVE_USER", "ONEDRIVE_FOLDER"]
GUID = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")

AADSTS_AYUDA = {
    "AADSTS7000215": "Secreto inválido: copiaste el 'Id del secreto' en vez del 'Valor', o el valor está mal pegado.",
    "AADSTS7000222": "El secreto EXPIRÓ. Crea uno nuevo en Azure (Certificados y secretos) y actualiza el Secret.",
    "AADSTS700016": "El Client ID no existe en ese tenant. Revisa AZURE_CLIENT_ID y AZURE_TENANT_ID.",
    "AADSTS90002": "Tenant no encontrado. Revisa AZURE_TENANT_ID.",
    "AADSTS900023": "Tenant mal escrito. Revisa AZURE_TENANT_ID.",
    "AADSTS53003": "Una política de acceso condicional bloquea la app. Avisa a TI.",
}


# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------
def secreto(clave: str) -> str:
    try:
        return str(st.secrets[clave]).strip()
    except Exception:
        return ""


def resumen_error(r: requests.Response) -> str:
    try:
        e = r.json().get("error", {})
        return f"HTTP {r.status_code} · {e.get('code', '?')} · {e.get('message', r.text[:200])}"
    except Exception:
        return f"HTTP {r.status_code} · {r.text[:200]}"


@st.cache_resource(show_spinner=False)
def _msal_app(tenant: str, client_id: str, secret: str):
    return msal.ConfidentialClientApplication(
        client_id, authority=f"https://login.microsoftonline.com/{tenant}", client_credential=secret
    )


def pedir_token() -> dict:
    app = _msal_app(secreto("AZURE_TENANT_ID"), secreto("AZURE_CLIENT_ID"), secreto("AZURE_CLIENT_SECRET"))
    return app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])


def decodificar_token(token: str) -> dict:
    payload = token.split(".")[1]
    payload += "=" * (-len(payload) % 4)
    return json.loads(base64.urlsafe_b64decode(payload))


def graph(metodo: str, ruta: str, token: str, **kw) -> requests.Response:
    url = ruta if ruta.startswith("http") else GRAPH + ruta
    headers = {"Authorization": f"Bearer {token}"}
    headers.update(kw.pop("headers", {}))
    return requests.request(metodo, url, headers=headers, timeout=60, **kw)


def base_drive() -> str:
    return f"/users/{requests.utils.quote(secreto('ONEDRIVE_USER'))}/drive"


def carpeta() -> str:
    return secreto("ONEDRIVE_FOLDER").strip("/")


def ruta_item(nombre: str) -> str:
    return f"{base_drive()}/root:/{requests.utils.quote(carpeta() + '/' + nombre, safe='/')}"


def es_marcador_footer(texto) -> bool:
    t = str(texto).upper()
    for ch in ("°", "º", " ", "\n"):
        t = t.replace(ch, "")
    return t.startswith("VB") and "CALIDAD" in t


class Pasos:
    """Acumula (estado, texto) donde estado es ok / error / aviso / info."""

    def __init__(self):
        self.items = []

    def ok(self, t): self.items.append(("ok", t))
    def error(self, t): self.items.append(("error", t))
    def aviso(self, t): self.items.append(("aviso", t))
    def info(self, t): self.items.append(("info", t))

    def si(self, condicion, texto_ok, texto_mal, nivel_mal="error"):
        """Agrega ok si se cumple la condición; si no, agrega el mensaje de fallo."""
        self.items.append(("ok", texto_ok) if condicion else (nivel_mal, texto_mal))

    @property
    def fallo(self): return any(e == "error" for e, _ in self.items)


ICONO = {"ok": "✅", "error": "❌", "aviso": "⚠️", "info": "ℹ️"}


def mostrar(nombre: str, pasos: Pasos):
    with st.container(border=True):
        st.markdown(f"**{'❌' if pasos.fallo else '✅'} {nombre}**")
        for estado, texto in pasos.items:
            st.markdown(f"{ICONO[estado]} {texto}")


# ----------------------------------------------------------------------------
# PRUEBA 1 - Secrets
# ----------------------------------------------------------------------------
def prueba_secrets() -> Pasos:
    p = Pasos()
    if not MSAL_OK:
        p.error("La librería `msal` no está instalada: revisa que `requirements.txt` la incluya.")
    for k in CLAVES:
        p.si(secreto(k), f"`{k}` presente", f"Falta `{k}` en los Secrets")
    cid, tid, sec, usr = (secreto(k) for k in ("AZURE_CLIENT_ID", "AZURE_TENANT_ID", "AZURE_CLIENT_SECRET", "ONEDRIVE_USER"))
    if cid and not GUID.match(cid): p.error("`AZURE_CLIENT_ID` no parece un GUID (formato xxxxxxxx-xxxx-...).")
    if tid and not GUID.match(tid): p.error("`AZURE_TENANT_ID` no parece un GUID.")
    if sec and GUID.match(sec):
        p.error("`AZURE_CLIENT_SECRET` parece un GUID: seguramente copiaste el **Id del secreto**. "
                "Debes usar la columna **Valor**.")
    elif sec and (len(sec) < 20 or "tu-client-secret" in sec):
        p.error("`AZURE_CLIENT_SECRET` parece incompleto o es el texto de ejemplo.")
    if usr and "@" not in usr: p.error("`ONEDRIVE_USER` debe ser un correo (UPN), ej. nombre@empresa.pe")
    if not p.fallo: p.info("Solo se comprueba que existan y tengan buen formato; nunca se muestran sus valores.")
    return p


# ----------------------------------------------------------------------------
# PRUEBA 2 - Token y permisos
# ----------------------------------------------------------------------------
def prueba_token() -> Pasos:
    p = Pasos()
    try:
        res = pedir_token()
    except Exception as e:
        p.error(f"No se pudo pedir el token: {type(e).__name__}: {str(e)[:200]}")
        return p
    if "access_token" not in res:
        desc = (res.get("error_description") or res.get("error") or "sin detalle").split("\r\n")[0][:300]
        p.error(f"Azure rechazó la autenticación: {desc}")
        for codigo, ayuda in AADSTS_AYUDA.items():
            if codigo in desc: p.aviso(ayuda)
        return p
    p.ok("Azure entregó un token de acceso.")
    try:
        claims = decodificar_token(res["access_token"])
        roles = claims.get("roles", [])
        exp = datetime.fromtimestamp(claims.get("exp", 0)).strftime("%H:%M:%S")
        p.info(f"Tenant: `{claims.get('tid', '?')}` · App: `{claims.get('appid', claims.get('azp', '?'))}` · vence {exp} (UTC del servidor)")
        if roles:
            p.ok("Permisos de aplicación aprobados: " + ", ".join(f"`{r}`" for r in roles))
            if not any(r in roles for r in ("Files.ReadWrite.All", "Sites.ReadWrite.All", "Sites.Selected")):
                p.error("Falta un permiso de archivos (`Files.ReadWrite.All` o `Sites.ReadWrite.All`).")
        else:
            p.error("El token NO trae permisos (`roles` vacío): falta el **consentimiento de administrador** "
                    "en Azure (Permisos de API → Conceder consentimiento). Pide a TI que lo apruebe.")
    except Exception:
        p.aviso("No pude leer los permisos del token, pero la autenticación funcionó.")
    return p


def token_o_none(p: Pasos):
    try:
        res = pedir_token()
        if "access_token" in res:
            return res["access_token"]
    except Exception:
        pass
    p.error("No se pudo obtener token (ejecuta antes la prueba 2).")
    return None


# ----------------------------------------------------------------------------
# PRUEBA 3 - OneDrive y carpeta
# ----------------------------------------------------------------------------
def prueba_onedrive() -> Pasos:
    p = Pasos()
    tk = token_o_none(p)
    if not tk: return p
    r = graph("GET", f"{base_drive()}?$select=id,name,driveType", tk)
    if r.status_code != 200:
        p.error("No se pudo abrir el OneDrive de ese usuario. " + resumen_error(r))
        if r.status_code == 403: p.aviso("403 = falta consentimiento de administrador o permiso `Files.ReadWrite.All`.")
        if r.status_code == 404: p.aviso("404 = el correo `ONEDRIVE_USER` no existe, o esa persona nunca abrió su OneDrive.")
        return p
    d = r.json()
    p.ok(f"OneDrive encontrado: **{d.get('name')}** (tipo `{d.get('driveType')}`)")
    r = graph("GET", f"{base_drive()}/root:/{requests.utils.quote(carpeta(), safe='/')}:/children?$select=name,size,lastModifiedDateTime&$top=100", tk)
    if r.status_code == 404:
        p.aviso(f"La carpeta `{carpeta()}` no existe todavía. Se creará sola al escribir el primer archivo (prueba 4), "
                "o créala tú en OneDrive.")
    elif r.status_code != 200:
        p.error("No se pudo listar la carpeta. " + resumen_error(r))
    else:
        items = r.json().get("value", [])
        p.ok(f"Carpeta `{carpeta()}` accesible · {len(items)} elemento(s).")
        if items:
            p.info("Contenido (máx. 100): " + ", ".join(f"`{i['name']}`" for i in items[:30]) + (" …" if len(items) > 30 else ""))
    return p


# ----------------------------------------------------------------------------
# PRUEBA 4 - Escribir / leer / borrar un archivo de texto
# ----------------------------------------------------------------------------
def prueba_archivo() -> Pasos:
    p = Pasos()
    tk = token_o_none(p)
    if not tk: return p
    nombre = f"{PREFIJO}_{int(time.time())}.txt"
    contenido = f"Prueba de vinculación {datetime.now().isoformat()}".encode()
    r = graph("PUT", ruta_item(nombre) + ":/content", tk, data=contenido, headers={"Content-Type": "text/plain"})
    if r.status_code not in (200, 201):
        p.error("No se pudo crear el archivo. " + resumen_error(r))
        if r.status_code == 403: p.aviso("403 = falta permiso de ESCRITURA (`Files.ReadWrite.All`, no solo `Files.Read.All`).")
        return p
    item_id = r.json()["id"]
    p.ok(f"Archivo creado: `{nombre}`")
    r = graph("GET", f"{base_drive()}/items/{item_id}/content", tk)
    p.si(r.status_code == 200 and r.content == contenido, "Archivo leído de vuelta y coincide.", "La lectura no coincide. " + resumen_error(r))
    r = graph("DELETE", f"{base_drive()}/items/{item_id}", tk)
    p.si(r.status_code == 204, "Archivo de prueba borrado.", "No se pudo borrar. " + resumen_error(r))
    return p


# ----------------------------------------------------------------------------
# PRUEBA 5 - Excel: lo que harán las apps reales
# ----------------------------------------------------------------------------
def _plantilla_xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja 1"
    ws["C1"] = "FORMATO DE PRUEBA"
    ws.merge_cells("C1:E2")
    for i, h in enumerate(["FECHA", "AREA", "CLIENTE", "N° de Muestra", "Producto"], 1):
        ws.cell(4, i, h)
    ws["A7"] = "V°B° Supervisor de Calidad"
    ws.merge_cells("A7:C7")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _agregar_registros(contenido: bytes) -> bytes:
    """Misma lógica que usarán las apps: copiar 'Hoja 1', limpiar filas de ejemplo e insertar
    las filas nuevas justo antes de la firma."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(contenido))
    ws = wb.copy_worksheet(wb["Hoja 1"])
    ws.title = "2026-01-01"
    footer = next(r for r in range(1, ws.max_row + 1) if es_marcador_footer(ws.cell(r, 1).value))
    if footer > 5:
        ws.delete_rows(5, footer - 5)
        footer = 5
    filas = [["01/01/2026", "EMPAQUE", "STARBUCKS", i, "PRODUCTO PRUEBA"] for i in (1, 2, 3)]
    ws.insert_rows(footer, len(filas))
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila, 1):
            ws.cell(footer + i, j, v)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def prueba_excel() -> Pasos:
    from openpyxl import load_workbook
    p = Pasos()
    tk = token_o_none(p)
    if not tk: return p
    nombre = f"{PREFIJO}_{int(time.time())}.xlsx"
    r = graph("PUT", ruta_item(nombre) + ":/content", tk, data=_plantilla_xlsx(),
              headers={"Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
    if r.status_code not in (200, 201):
        p.error("No se pudo subir el Excel de prueba. " + resumen_error(r))
        return p
    item_id = r.json()["id"]
    p.ok(f"Excel de prueba subido: `{nombre}`")
    try:
        # (a) API 'Excel' de Microsoft con permisos de aplicación (informativo)
        r = graph("POST", f"{base_drive()}/items/{item_id}/workbook/createSession", tk, json={"persistChanges": True})
        if r.status_code in (200, 201):
            p.info("La API de Excel de Graph (`/workbook`) SÍ funcionó con permisos de aplicación en tu tenant.")
        else:
            p.info("La API de Excel de Graph (`/workbook`) no acepta permisos de aplicación (esperado según la "
                   "documentación de Microsoft). Por eso las apps editan el archivo directamente. "
                   f"[{resumen_error(r)[:120]}]")
        # (b) Descargar -> editar con openpyxl -> subir con control de versión (eTag)
        meta = graph("GET", f"{base_drive()}/items/{item_id}?$select=eTag", tk).json()
        etag = meta.get("eTag")
        original = graph("GET", f"{base_drive()}/items/{item_id}/content", tk).content
        nuevo = _agregar_registros(original)
        p.ok("Excel descargado y editado (hoja copiada + 3 filas insertadas antes de la firma).")
        r = graph("PUT", f"{base_drive()}/items/{item_id}/content", tk, data=nuevo,
                  headers={"If-Match": etag, "Content-Type": "application/octet-stream"})
        if r.status_code not in (200, 201):
            p.error("No se pudo subir el Excel editado. " + resumen_error(r))
            return p
        p.ok("Excel editado subido de vuelta (con `If-Match`).")
        # verificar
        wb = load_workbook(io.BytesIO(graph("GET", f"{base_drive()}/items/{item_id}/content", tk).content))
        ws = wb["2026-01-01"]
        ok_filas = [ws.cell(r, 4).value for r in (5, 6, 7)] == [1, 2, 3] and es_marcador_footer(ws.cell(8, 1).value)
        p.si(ok_filas, "Verificado: 3 filas quedaron en 5–7 y la firma en la fila 8.", "Las filas no quedaron donde se esperaba.")
        # (c) protección contra pisar cambios: subir con un eTag viejo debe fallar (412)
        r = graph("PUT", f"{base_drive()}/items/{item_id}/content", tk, data=nuevo,
                  headers={"If-Match": etag, "Content-Type": "application/octet-stream"})
        if r.status_code == 412:
            p.ok("Protección contra guardados simultáneos OK (con un eTag viejo Microsoft rechazó con 412).")
        else:
            p.aviso(f"Con un eTag viejo esperaba 412 y recibí {r.status_code}: el control de versión no se aplicó.")
    except Exception as e:
        p.error(f"Error durante la prueba de Excel: {type(e).__name__}: {str(e)[:200]}")
    finally:
        r = graph("DELETE", f"{base_drive()}/items/{item_id}", tk)
        p.si(r.status_code == 204, "Excel de prueba borrado.", "No se pudo borrar el Excel de prueba: bórralo a mano en OneDrive.", "aviso")
    return p


def limpiar() -> Pasos:
    p = Pasos()
    tk = token_o_none(p)
    if not tk: return p
    r = graph("GET", f"{base_drive()}/root:/{requests.utils.quote(carpeta(), safe='/')}:/children?$select=id,name&$top=200", tk)
    if r.status_code != 200:
        p.aviso("No pude listar la carpeta: " + resumen_error(r))
        return p
    borrados = 0
    for it in r.json().get("value", []):
        if it["name"].startswith(PREFIJO):
            if graph("DELETE", f"{base_drive()}/items/{it['id']}", tk).status_code == 204: borrados += 1
    p.ok(f"Archivos de prueba borrados: {borrados}")
    return p


# ----------------------------------------------------------------------------
# INTERFAZ
# ----------------------------------------------------------------------------
PIN = secreto("APP_PIN")
if PIN and st.session_state.get("pin_ok") is not True:
    st.title("🔗 Prueba de vinculación")
    ingreso = st.text_input("PIN de acceso", type="password")
    if st.button("Entrar"):
        if ingreso == PIN:
            st.session_state.pin_ok = True
            st.rerun()
        else:
            st.error("PIN incorrecto.")
    st.stop()

st.title("🔗 Prueba de vinculación con OneDrive")
st.caption("App independiente: no toca tus apps de producción. Solo crea y borra archivos "
           f"`{PREFIJO}*` en la carpeta configurada.")

if "resultados" not in st.session_state:
    st.session_state.resultados = {}

PRUEBAS = [
    ("1. Secrets configurados", prueba_secrets),
    ("2. Token y permisos de Azure", prueba_token),
    ("3. OneDrive y carpeta de trabajo", prueba_onedrive),
    ("4. Crear, leer y borrar un archivo", prueba_archivo),
    ("5. Editar un Excel como lo harán las apps", prueba_excel),
]

if st.button("▶️ Ejecutar todas las pruebas", type="primary"):
    st.session_state.resultados = {}
    for nombre, fn in PRUEBAS:
        with st.spinner(nombre):
            st.session_state.resultados[nombre] = fn()
        if st.session_state.resultados[nombre].fallo and nombre.startswith(("1.", "2.")):
            break  # sin secrets o token válidos no tiene sentido seguir

st.markdown("**O una por una:**")
cols = st.columns(len(PRUEBAS))
for col, (nombre, fn) in zip(cols, PRUEBAS):
    if col.button(nombre.split(".")[0], help=nombre):
        with st.spinner(nombre):
            st.session_state.resultados[nombre] = fn()

for nombre, _ in PRUEBAS:
    if nombre in st.session_state.resultados:
        mostrar(nombre, st.session_state.resultados[nombre])

st.divider()
if st.button("🧹 Borrar archivos de prueba que hayan quedado"):
    mostrar("Limpieza", limpiar())

if st.session_state.resultados and len(st.session_state.resultados) == len(PRUEBAS):
    if not any(r.fallo for r in st.session_state.resultados.values()):
        st.success("🎉 Vinculación lista: ya se puede migrar una app real a OneDrive.")
